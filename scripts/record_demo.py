"""Record the 2-3 minute demo video automatically (follows docs/demo-video-script.md).

Needs:  pip install playwright imageio-ffmpeg && playwright install chromium
Run the app first (./run.sh), then:
    python -m scripts.record_demo --url http://localhost:8001
Output: docs/demo.mp4 (and static/media/demo.webm for the landing page)

Add your voice-over afterwards, or use the on-screen captions as they are.
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import tempfile
from pathlib import Path

from playwright.sync_api import sync_playwright

BASE = Path(__file__).resolve().parent.parent
SAMPLES = BASE / "sample_data" / "invoices"
FILES = ["alnoor_INV-2041.pdf", "crescent_CP-26-0187.pdf", "photo_alnoor_INV-2063.jpg", "scan_margalla_ME-5544.pdf",
         "indus_office_IOS-1204_totals_wrong.pdf", "crescent_invoice_copy_from_email.pdf", "margalla_ME-5521.pdf",
         "barkat_karyana_BK-0415_urdu.pdf", "swift_courier_SC-88213.pdf", "global_imports_GTI-7781.pdf"]

CAPTION_JS = """
(text) => {
  let el = document.getElementById('__cap');
  if (!el) {
    el = document.createElement('div'); el.id = '__cap';
    el.style.cssText = 'position:fixed;left:50%;bottom:36px;transform:translateX(-50%);z-index:99999;max-width:980px;' +
      'background:rgba(17,24,39,.92);color:#fff;font:600 24px/1.35 Inter,system-ui,sans-serif;padding:16px 26px;border-radius:14px;' +
      'box-shadow:0 10px 40px rgba(0,0,0,.35);text-align:center;transition:opacity .25s;pointer-events:none';
    document.body.appendChild(el);
  }
  el.style.opacity = text ? '1' : '0';
  if (text) el.innerHTML = text;
}
"""


def caption(page, text, ms=3000):
    page.evaluate(CAPTION_JS, text)
    page.wait_for_timeout(ms)


def excel_preview_html(xlsx_path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path)
    tabs = "".join(f'<span class="tab {"on" if i == 0 else ""}">{ws.title}</span>' for i, ws in enumerate(wb.worksheets))
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))[:12]
    def cell(v):
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return f"{v:,.2f}"
        if hasattr(v, "strftime"):
            return v.strftime("%d/%m/%Y")
        return "" if v is None else str(v)
    head = "".join(f"<th>{cell(v)}</th>" for v in rows[0])
    body = "".join("<tr>" + "".join(f"<td>{cell(v)}</td>" for v in r) + "</tr>" for r in rows[1:])
    return f"""<html><body style="margin:0;font-family:Inter,Calibri,Arial;background:#fff">
      <div style="background:#107c41;color:#fff;padding:10px 18px;font:600 15px Arial">DocuFlow export.xlsx - Excel</div>
      <div style="overflow:hidden;padding:0"><table style="border-collapse:collapse;font-size:13px">
      <style>th{{background:#1e3a8a;color:#fff;padding:7px 10px;white-space:nowrap;border:1px solid #d9dee7}} td{{padding:6px 10px;border:1px solid #e5e7eb;white-space:nowrap}}
      .tab{{display:inline-block;padding:6px 16px;border:1px solid #ccc;border-top:0;background:#f3f3f3;font-size:13px}} .tab.on{{background:#fff;color:#107c41;font-weight:700}}</style>
      <tr>{head}</tr>{body}</table></div>
      <div style="position:fixed;bottom:0;left:0;right:0;background:#f3f3f3;border-top:1px solid #ccc">{tabs}</div></body></html>"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--url", default="http://localhost:8001")
    args = ap.parse_args()
    tmp = Path(tempfile.mkdtemp())
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, channel="chrome") if shutil.which("google-chrome") else p.chromium.launch(headless=True)
        ctx = browser.new_context(viewport={"width": 1280, "height": 720}, record_video_dir=str(tmp),
                                  record_video_size={"width": 1280, "height": 720}, accept_downloads=True)
        page = ctx.new_page()
        page.goto(args.url + "/#/login")
        page.wait_for_timeout(1200)
        caption(page, "How long does your team spend typing invoices every month?", 3800)
        caption(page, "DocuFlow turns invoices, bills and forms into Excel. No typing.", 3000)
        page.click("#demoBtn")
        page.wait_for_timeout(1800)
        page.goto(args.url + "/#/")
        page.wait_for_timeout(1200)
        caption(page, "Drop in a mix: clean PDFs, a scanned copy, phone photos, an Urdu-English bill", 2500)
        page.set_input_files("#fileInput", [str(SAMPLES / f) for f in FILES])
        page.wait_for_timeout(1500)
        caption(page, "", 100)
        page.click("#startBtn")
        page.wait_for_url("**/processing", timeout=15000)
        caption(page, "Every file is read, then every value is checked", 1500)
        page.wait_for_url(lambda u: "/processing" not in u, timeout=180000)
        page.wait_for_timeout(1500)
        caption(page, "Files processed, fields extracted, what needs attention, and time saved", 4200)
        caption(page, "", 100)
        page.click("text=Review")
        page.wait_for_timeout(1800)

        for _ in range(6):
            fname = page.inner_text(".viewer-bar .fname")
            if "scan_" in fname:
                caption(page, "Orange means \"please check\". Here the customer name was hard to read on the scan", 3800)
                page.click("#f_customer")
                page.wait_for_timeout(500)
                page.type("#f_customer", "Bright Mart Superstore", delay=45)
                caption(page, "Fix it in a second, then press Enter to approve. The next file opens automatically", 3200)
                caption(page, "", 100)
                page.keyboard.press("Enter")
            elif "totals_wrong" in fname:
                page.click("#f_total")
                caption(page, "Totals don't add up? Caught before it reaches your books", 3800)
                caption(page, "Click any value to see exactly where it came from on the page", 3000)
                caption(page, "", 100)
                page.click("#flagBtn")
            elif "copy_from_email" in fname:
                caption(page, "Same invoice received twice? Flagged as a possible duplicate", 3800)
                caption(page, "", 100)
                page.click("#skipBtn")
            else:
                page.keyboard.press("Enter")
            page.wait_for_timeout(1800)
            if "/review/" not in page.url:
                break

        page.wait_for_timeout(1200)
        caption(page, "Download Excel: Invoices, Line Items and Needs Review sheets", 2500)
        with page.expect_download() as dl:
            page.click("#dlMain")
        xlsx = tmp / "export.xlsx"
        dl.value.save_as(str(xlsx))
        caption(page, "", 100)
        page.set_content(excel_preview_html(xlsx))
        page.wait_for_timeout(600)
        caption(page, "Opens straight in Excel or Google Sheets, in your own column names", 3800)
        results = BASE / "docs" / "results.json"
        if results.exists():
            import json
            s = json.loads(results.read_text())["summary"]
            caption(page, f"On our {s['documents']}-document test set: {s['field_accuracy_readable']}% of fields correct on readable documents,<br>"
                          f"{s['avg_seconds_per_doc']} seconds per document, wrong totals and duplicates caught", 5200)
        caption(page, "Try it free on 20 of your own invoices", 3800)
        video = page.video.path()
        ctx.close()
        browser.close()

    out_webm = BASE / "static" / "media" / "demo.webm"
    out_webm.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(video, out_webm)
    try:
        import imageio_ffmpeg
        subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-y", "-loglevel", "error", "-i", str(video), "-c:v", "libx264",
                        "-pix_fmt", "yuv420p", "-crf", "23", "-movflags", "+faststart", str(BASE / "docs" / "demo.mp4")], check=True)
        print("Saved docs/demo.mp4 and static/media/demo.webm")
    except Exception as e:
        print("Saved static/media/demo.webm (mp4 conversion skipped:", e, ")")


if __name__ == "__main__":
    main()
