"""Excel / CSV output in the client's own column names and date format."""
from __future__ import annotations

import csv
import io
import json
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .presets import FIELDS, default_config

DATE_FORMATS = {"DD/MM/YYYY": ("%d/%m/%Y", "DD/MM/YYYY"), "YYYY-MM-DD": ("%Y-%m-%d", "YYYY-MM-DD"),
                "DD Mon YYYY": ("%d %b %Y", "DD MMM YYYY"), "MM/DD/YYYY": ("%m/%d/%Y", "MM/DD/YYYY")}
STATUS_LABEL = {"approved": "Approved", "done": "Ready", "needs_review": "Needs review", "skipped": "Skipped",
                "flagged": "Flagged", "failed": "Could not read"}


def _columns(doc_type, config):
    config = config or default_config(doc_type)
    kinds = {k: kind for k, _, kind in FIELDS.get(doc_type, [])}
    cols = [(f["key"], f["label"], kinds.get(f["key"], "text")) for f in config["fields"] if f.get("enabled", True)]
    return cols, config


def _rows(docs, doc_type, config):
    cols, config = _columns(doc_type, config)
    if doc_type == "form":
        keys = []
        for d in docs:
            for k in json.loads(d["fields_json"] or "{}"):
                if k not in keys:
                    keys.append(k)
        cols = [(k, k, "text") for k in keys]
    return cols, config


def _to_date(v):
    try:
        return date.fromisoformat(str(v))
    except (TypeError, ValueError):
        return None


def build_xlsx(batch, docs, config) -> bytes:
    doc_type = batch["doc_type"]
    cols, config = _rows(docs, doc_type, config)
    date_py, date_xl = DATE_FORMATS.get(config.get("date_format", "DD/MM/YYYY"), DATE_FORMATS["DD/MM/YYYY"])
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9DEE7")
    wb = Workbook()

    def sheet(ws, headers, rows, kinds):
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 22
        for r in rows:
            ws.append(r)
        for ci, kind in enumerate(kinds, start=1):
            letter = get_column_letter(ci)
            width = max([len(str(headers[ci - 1]))] + [len(str(r[ci - 1])) for r in rows if r[ci - 1] is not None] + [8])
            ws.column_dimensions[letter].width = min(width + 3, 48)
            for cell in ws[letter][1:]:
                cell.border = Border(bottom=thin)
                if kind == "money":
                    cell.number_format = "#,##0.00"
                elif kind == "date" and isinstance(cell.value, date):
                    cell.number_format = date_xl
                elif kind == "int":
                    cell.number_format = "#,##0.##"
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions

    # Sheet 1: one row per document
    ws = wb.active
    ws.title = "Invoices" if doc_type != "form" else "Forms"
    headers = ["File"] + [label for _, label, _ in cols] + ["Status"]
    kinds = ["text"] + [k for _, _, k in cols] + ["text"]
    rows = []
    for d in docs:
        f = json.loads(d["fields_json"] or "{}")
        row = [d["filename"]]
        for key, _, kind in cols:
            v = f.get(key)
            row.append(_to_date(v) or v if kind == "date" else v)
        row.append(STATUS_LABEL.get(d["status"], d["status"]))
        rows.append(row)
    sheet(ws, headers, rows, kinds)

    # Sheet 2: line items linked by invoice number
    if doc_type != "form" and config.get("include_line_items", True):
        ws2 = wb.create_sheet("Line Items")
        inv_label = next((label for key, label, _ in cols if key == "invoice_no"), "Invoice No")
        rows2 = []
        for d in docs:
            f = json.loads(d["fields_json"] or "{}")
            for n, it in enumerate(json.loads(d["items_json"] or "[]"), start=1):
                rows2.append([f.get("invoice_no") or d["filename"], f.get("vendor"), n, it.get("description"),
                              it.get("quantity"), it.get("unit_price"), it.get("amount")])
        sheet(ws2, [inv_label, "Vendor", "Line", "Description", "Qty", "Unit Price", "Amount"], rows2,
              ["text", "text", "int", "text", "int", "money", "money"])

    # Sheet 3: anything that still needs a person
    ws3 = wb.create_sheet("Needs Review")
    rows3 = []
    for d in docs:
        reasons = []
        checks = json.loads(d["checks_json"] or "{}")
        labels = {k: label for k, label, _ in cols}
        for k, c in checks.items():
            if isinstance(c, dict) and c.get("status") in ("warn", "error") and d["status"] not in ("approved",):
                reasons.append(f"{labels.get(k, k)}: {c.get('reason')}")
        for w in json.loads(d["warnings_json"] or "[]"):
            if d["status"] != "approved" or w["type"] == "duplicate":
                reasons.append(w["text"])
        if d["status"] == "failed":
            reasons.append(d.get("error") or "Could not read")
        if d["status"] == "flagged":
            reasons.insert(0, "Flagged by reviewer")
        if reasons:
            f = json.loads(d["fields_json"] or "{}")
            rows3.append([d["filename"], f.get("invoice_no"), f.get("vendor"), STATUS_LABEL.get(d["status"], d["status"]), " | ".join(reasons)])
    sheet(ws3, ["File", "Invoice No", "Vendor", "Status", "Reason"], rows3, ["text"] * 5)
    ws3.column_dimensions["E"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _n(v):
    """Whole numbers without '.0' in CSV."""
    return int(v) if isinstance(v, float) and v.is_integer() else v


def build_csv(batch, docs, config, sheet="invoices") -> bytes:
    doc_type = batch["doc_type"]
    cols, config = _rows(docs, doc_type, config)
    date_py, _ = DATE_FORMATS.get(config.get("date_format", "DD/MM/YYYY"), DATE_FORMATS["DD/MM/YYYY"])
    buf = io.StringIO()
    w = csv.writer(buf)
    if sheet == "line_items":
        w.writerow(["Invoice No", "Vendor", "Line", "Description", "Qty", "Unit Price", "Amount"])
        for d in docs:
            f = json.loads(d["fields_json"] or "{}")
            for n, it in enumerate(json.loads(d["items_json"] or "[]"), start=1):
                w.writerow([f.get("invoice_no") or d["filename"], f.get("vendor"), n, it.get("description"),
                            _n(it.get("quantity")), it.get("unit_price"), it.get("amount")])
    else:
        w.writerow(["File"] + [label for _, label, _ in cols] + ["Status"])
        for d in docs:
            f = json.loads(d["fields_json"] or "{}")
            row = [d["filename"]]
            for key, _, kind in cols:
                v = f.get(key)
                if kind == "date" and _to_date(v):
                    v = _to_date(v).strftime(date_py)
                row.append("" if v is None else v)
            row.append(STATUS_LABEL.get(d["status"], d["status"]))
            w.writerow(row)
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 (Urdu) correctly
